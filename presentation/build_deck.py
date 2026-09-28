"""Generate the Arabic, nine-slide Qaloon hackathon deck (PPTX and PDF).

Requires: pillow, python-pptx, reportlab, arabic-reshaper, python-bidi.
Uses the installed Windows Tahoma typeface for consistent Arabic shaping.
"""

from io import BytesIO
from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as pdf_canvas


HERE = Path(__file__).resolve().parent
W, H = 1600, 900
NAVY = "#101C2A"
DEEP = "#162536"
CARD = "#203245"
CREAM = "#F6F3E9"
PALE = "#C5D2D0"
MINT = "#7DD9C0"
GOLD = "#E6BE83"
CORAL = "#F3A995"
FONT_REG = "C:/Windows/Fonts/tahoma.ttf"
FONT_BOLD = "C:/Windows/Fonts/tahomabd.ttf"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def rtl(value):
    return get_display(arabic_reshaper.reshape(str(value)))


def text(draw, value, right, top, size=30, color=CREAM, bold=False):
    value = rtl(value)
    draw.text((right, top), value, font=font(size, bold), fill=color, anchor="ra")


def left(draw, value, x, top, size=30, color=CREAM, bold=False):
    draw.text((x, top), value, font=font(size, bold), fill=color, anchor="la")


def paragraph(draw, value, right, top, width, size=27, color=PALE,
              spacing=17, bold=False):
    words = value.split()
    lines, line = [], ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if line and draw.textlength(rtl(candidate), font=font(size, bold)) > width:
            lines.append(line)
            line = word
        else:
            line = candidate
    if line:
        lines.append(line)
    for index, item in enumerate(lines):
        text(draw, item, right, top + index * (size + spacing), size, color, bold)
    return top + len(lines) * (size + spacing)


def round_rect(draw, xy, fill=CARD, outline=None, radius=24, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def base(number, section, title=None, subtitle=None):
    im = Image.new("RGB", (W, H), NAVY)
    d = ImageDraw.Draw(im)
    # Restrained geometric accent; no stock imagery or ornamental clutter.
    d.ellipse((-220, 630, 320, 1170), outline="#244350", width=2)
    d.ellipse((-140, 710, 240, 1090), outline="#244350", width=2)
    d.rectangle((75, 75, 145, 81), fill=MINT)
    left(d, "QĀLŪN / ASR", 75, 51, 19, MINT, True)
    text(d, section, 1525, 53, 21, PALE)
    d.line((75, 833, 1525, 833), fill="#314555", width=2)
    text(d, "مقرأة قالون  /  عرض التحدّي", 410, 848, 18, PALE)
    left(d, f"{number:02d} / 09", 1410, 848, 18, MINT, True)
    if title:
        text(d, title, 1520, 123, 57, CREAM, True)
    if subtitle:
        text(d, subtitle, 1518, 204, 25, PALE)
    return im, d


def pill(d, label, x, y, width, color=MINT):
    round_rect(d, (x, y, x + width, y + 53), fill="#294448", radius=25)
    text(d, label, x + width - 18, y + 12, 21, color, True)


def tile(d, box, label, heading, body, accent=MINT):
    x1, y1, x2, y2 = box
    round_rect(d, box)
    d.rectangle((x2 - 9, y1 + 25, x2 - 4, y2 - 25), fill=accent)
    text(d, label, x2 - 36, y1 + 28, 21, accent, True)
    text(d, heading, x2 - 36, y1 + 84, 36, CREAM, True)
    paragraph(d, body, x2 - 36, y1 + 153, x2 - x1 - 80, 26)


slides = []

# 1 — The complete intended experience.
im, d = base(1, "التجارب التفاعلية والرحلة المعرفية")
round_rect(d, (100, 250, 1500, 752), fill=DEEP, radius=36)
d.rectangle((1462, 298, 1472, 696), fill=MINT)
text(d, "مقرأة قالون", 1395, 312, 93, CREAM, True)
text(d, "تطبيق تسميع ذكي لرواية قالون", 1390, 449, 43, MINT)
paragraph(d, "اختر آية من القرآن الكريم، واتلُها؛ يتابعك المساعد من هناك، ويعينك على تذكّر الكلمات المنسية ومواصلة الحفظ.",
          1390, 548, 1050, 31, CREAM, 20)
pill(d, "القرآن الكريم كاملاً  •  قالون", 1080, 687, 320)
slides.append(im)

# 2 — The gap: no unsupportable prevalence claims.
im, d = base(2, "01  /  الفجوة", "صوت قالون يحتاج من يفهمه", "التجربة الصوتية الشائعة لا تراعي دائماً خصوصية رواية قالون.")
tile(d, (820, 300, 1520, 713), "عند التلاوة", "اختلاف الرواية", "قد تُقارن قراءة صحيحة بنص رواية أخرى، فلا يحصل الحافظ على متابعة تناسب ما يتلوه.")
tile(d, (78, 300, 775, 713), "عند البناء", "غياب البيانات الجاهزة", "لا تتوفر مجموعة مفتوحة وشاملة لتسجيلات قالون المقطّعة آيةً آيةً كما يحتاجها التطبيق.", GOLD)
slides.append(im)

# 3 — Human interaction, not speculative performance.
im, d = base(3, "02  /  تجربة المتعلّم", "اختر الآية… ودع المقرأة تتابعك", "تجربة تسميع بسيطة، شبيهة بما يعرفه المستخدم؛ ولكن برواية قالون.")
for box, step, heading, body, accent in [
    ((1050, 305, 1520, 710), "01", "يختار", "يحدد الآية التي يريد حفظها أو مراجعتها فقط.", MINT),
    ((565, 305, 1035, 710), "02", "يتلو", "يستمع التطبيق ويتابع التلاوة من موضع الاختيار.", GOLD),
    ((80, 305, 550, 710), "03", "يتذكّر", "ينبّهه للكلمات المنسية ويتيح إعادة المحاولة.", CORAL),
]:
    tile(d, box, step, heading, body, accent)
slides.append(im)

# 4 — Plain-language explanation.
im, d = base(4, "03  /  كيف يعمل", "تجربة كاملة، لا مجرد تفريغ صوتي", "من اختيار الآية إلى الملاحظة التي تساعد الحافظ على الاستمرار.")
for y, number, heading, body, accent in [
    (283, "01", "يستمع للتلاوة", "يربط صوت المستخدم بالآية التي اختارها، ويتابع تقدّمه في القراءة.", MINT),
    (408, "02", "يراعي رواية قالون", "يقارن بما يوافق نص الرواية، لا بنص مختلف عنها.", GOLD),
    (533, "03", "يساعد على التذكّر", "يُظهر موضع الكلمة المنسية ويتيح للقارئ مواصلة التسميع.", CORAL),
]:
    round_rect(d, (120, y, 1480, y + 105), fill=DEEP, radius=19)
    left(d, number, 160, y + 20, 38, accent, True)
    text(d, heading, 1400, y + 10, 33, CREAM, True)
    text(d, body, 1400, y + 60, 23, PALE)
text(d, "الهدف: إرشاد عملي للحافظ مع بقاء المعلّم مرجعاً للتصحيح.", 1450, 721, 23, MINT)
slides.append(im)

# 5 — Data collection as a core part of the full-Quran vision.
im, d = base(5, "04  /  تأسيس البيانات", "سنبني بيانات قالون بأنفسنا", "غياب مصدر مفتوح شامل للآيات الصوتية يجعل جمع البيانات جزءاً من الحل.")
tile(d, (820, 300, 1520, 713), "نجمع", "التلاوات والنص", "نجمع تسجيلات قالون ونصوصه المعتمدة من مصادر منشورة على الإنترنت، مع مراعاة حقوق استخدامها.", MINT)
tile(d, (78, 300, 775, 713), "نُعِدّ", "القرآن كاملاً", "نراجع مطابقة الآيات صوتياً ونصياً، ونهيّئ بيانات صالحة لتدريب المساعد على القرآن كله.", GOLD)
slides.append(im)

# 6 — Product vision, without overclaiming a finished model.
im, d = base(6, "05  /  المنتج", "مقرأة رقمية في رحلة واحدة", "الحل المستهدف يعمل من البداية إلى النهاية على كامل القرآن برواية قالون.")
tile(d, (815, 310, 1520, 710), "للحافظ", "تسميع ومراجعة", "يبدأ من الآية التي اختارها، ويستمع إلى ملاحظة واضحة إذا نسي كلمة أو غيّر ترتيبها.", MINT)
tile(d, (80, 310, 785, 710), "للمعلّم", "متابعة مساندة", "يستخدم المساعد لتكرار المراجعة بين الجلسات، مع بقاء الحكم على جودة التلاوة بيد المعلّم.", GOLD)
slides.append(im)

# 7 — Evaluation readable for nontechnical judges.
im, d = base(7, "06  /  الجودة", "لا يكفي أن يسمع؛ يجب أن يفهم", "سنختبر جودة التجربة مع قرّاء ومعلّمين قبل نشرها على نطاق واسع.")
tile(d, (1035, 307, 1520, 703), "فهم التلاوة", "دقة الكلمات", "هل يكتب التطبيق ما تلاه القارئ فعلاً؟", MINT)
tile(d, (543, 307, 1020, 703), "التنبيه", "الكلمات المنسية", "هل يحدد موضع النسيان دون إنذارات خاطئة؟", GOLD)
tile(d, (70, 307, 528, 703), "التجربة", "رأي المعلّمين", "هل تساعد الملاحظات على الحفظ دون تشتيت؟", CORAL)
text(d, "لا ندّعي نسبة دقة قبل الاختبار؛ سننشر النتائج مع حدودها.", 1470, 748, 22, PALE)
slides.append(im)

# 8 — Full-Quran main release; Tajweed and articulation in V2.
im, d = base(8, "07  /  خارطة الطريق", "القرآن كاملاً أولاً… ثم فهمٌ أعمق للتجويد", "الإصدار الأساسي منتج متكامل، والإصدار الثاني يضيف مهارات تلاوة متقدمة.")
for y, phase, heading, description in [
    (300, "المرحلة ١", "جمع القرآن كاملاً", "تجميع التلاوات والنصوص ومراجعتها برواية قالون."),
    (420, "الإصدار ١", "التسميع من البداية للنهاية", "اختيار الآية، متابعة الحفظ، والتنبيه للكلمات المنسية."),
    (540, "الإصدار ٢", "التجويد ومخارج الحروف", "بحث تنبيهات نطقية تراعي أحكام الرواية وتخضع لمراجعة معلّمين."),
]:
    round_rect(d, (122, y, 1480, y + 98), fill=DEEP, radius=20)
    pill(d, phase, 165, y + 21, 165)
    text(d, heading, 1394, y + 11, 32, CREAM, True)
    text(d, description, 1394, y + 59, 23, PALE)
slides.append(im)

# 9 — Open source and impact.
im, d = base(9, "08  /  الأثر", "قالون للجميع… والأجر للجميع", "أداة تعليمية تُكمل دور المعلّم، ولا تستبدله.")
round_rect(d, (125, 290, 1475, 708), fill=DEEP, radius=34)
text(d, "منفعة تتجاوز حدود المسابقة", 1385, 330, 28, MINT, True)
left(d, "HUGGING FACE  /  OPEN SOURCE", 205, 336, 22, MINT, True)
paragraph(d, "نطمح لخدمة نحو 50 مليون قارئ ومتعلم لرواية قالون، ونشر النموذج والموارد المتاح نشرها مفتوحة المصدر؛ علماً نافعاً وأجراً مستمراً.",
          1380, 423, 1110, 37, CREAM, 18, True)
text(d, "نحتاج شركاء من المعلّمين لتجربة المنتج ومراجعة جودة التلاوة.", 1380, 625, 25, GOLD)
slides.append(im)

assert len(slides) == 9
pptx = Presentation()
pptx.slide_width, pptx.slide_height = Inches(13.333333), Inches(7.5)
pdf = pdf_canvas.Canvas(str(HERE / "Qaloon_Hackathon.pdf"), pagesize=(960, 540), pageCompression=1)
for im in slides:
    content = BytesIO()
    im.save(content, format="PNG", optimize=True)
    content.seek(0)
    pptx.slides.add_slide(pptx.slide_layouts[6]).shapes.add_picture(
        content, 0, 0, width=pptx.slide_width, height=pptx.slide_height
    )
    content.seek(0)
    pdf.drawImage(ImageReader(content), 0, 0, width=960, height=540)
    pdf.showPage()
pdf.save()
pptx.save(HERE / "Qaloon_Hackathon.pptx")
print("Created 9 slides: presentation/Qaloon_Hackathon.pptx and .pdf")
